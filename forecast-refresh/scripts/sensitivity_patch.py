#!/usr/bin/env python3
"""Fill the Readout 'RECOMMIT SENSITIVITY' block: committed-365d net if the recommit lands at
{no recommit / 25% / 30% / 35%}. Computes from the workbook's own ledger + Math baseline + list_rate.
Usage: sensitivity_patch.py --xlsx WB"""
import argparse
from datetime import date, datetime, timedelta
import openpyxl

def asdate(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    if isinstance(v,str):
        for f in ("%Y-%m-%d","%m/%d/%Y"):
            try: return datetime.strptime(v.strip(),f).date()
            except ValueError: pass
    return None

def lblcell(ws,key):
    for r in range(1,20):
        for c in range(1,8):
            if key in str(ws.cell(r,c).value or "").lower(): return ws.cell(r,c+2).value
    return None

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--xlsx",required=True); a=ap.parse_args()
    wb=openpyxl.load_workbook(a.xlsx, data_only=True)
    L=wb["Event Ledger"]
    list_rate=float(lblcell(L,"list_rate") or 0.52); as_of=asdate(lblcell(L,"as_of"))
    # ledger
    hr=next(r for r in range(1,30) if str(L.cell(r,2).value or "").strip().lower()=="event_id")
    contracts=[]; workloads=[]; recommit_date=None
    r=hr+1
    while L.cell(r,2).value:
        dt=asdate(L.cell(r,3).value); typ=str(L.cell(r,4).value or "").lower()
        stat=str(L.cell(r,5).value or "").lower(); pv=L.cell(r,7).value; inl=str(L.cell(r,9).value or "").upper()
        if typ=="contract" and dt:
            contracts.append((dt,stat,float(pv)))
            if stat=="projected" and recommit_date is None: recommit_date=dt
        if typ=="workload" and dt: workloads.append((dt,float(pv),inl))
        r+=1
    contracts.sort()
    current_disc=0.30
    for dt,stat,pv in contracts:
        if stat=="occurred" and dt<=as_of: current_disc=pv   # rate in force right now
    def wload(dt): return sum(p for wd,p,inl in workloads if inl=="Y" and wd<=dt)
    # Math: baseline DBU/day = col E(5); days = col D(4)  (daily actuals + weekly forecast)
    M=wb["Math"]; mhr=next(r for r in range(1,20) if str(M.cell(r,2).value or "").strip().lower()=="date")
    base={}; daysmap={}; dates=[]
    r=mhr+1
    while M.cell(r,2).value:
        dt=asdate(M.cell(r,2).value)
        if dt: dates.append(dt); base[dt]=float(M.cell(r,5).value or 0); daysmap[dt]=float(M.cell(r,4).value or 1)
        r+=1
    fut=[d for d in dates if as_of<d<=as_of+timedelta(days=365)]
    def committed(recommit_disc):
        s=0.0
        for d in fut:
            disc = current_disc if (recommit_date is None or d<recommit_date or recommit_disc is None) else recommit_disc
            s += (base.get(d,0)+wload(d))*daysmap.get(d,1)*list_rate*(1-disc)
        return s
    vals={"no":committed(None),"25":committed(0.25),"30":committed(0.30),"35":committed(0.35)}

    # write into Readout: find 'RECOMMIT SENSITIVITY' label, values 2 rows below at cols B,E,G,I
    wb2=openpyxl.load_workbook(a.xlsx)  # non-data to preserve formulas
    R=wb2["Readout"]; srow=None
    for r in range(1,80):
        if "recommit sensitivity" in str(R.cell(r,2).value or "").lower(): srow=r+2; break
    if srow is None: print("SENS label not found"); return
    for col,key in [(2,"no"),(5,"25"),(7,"30"),(9,"35")]:
        R.cell(srow,col, round(vals[key],0)).number_format="$#,##0"
    try: wb2.calculation.fullCalcOnLoad=True
    except Exception: pass
    wb2.save(a.xlsx)
    print(f"sensitivity: no-recommit ${vals['no']:,.0f} | @25% ${vals['25']:,.0f} | @30% ${vals['30']:,.0f} | @35% ${vals['35']:,.0f}  (current_disc {current_disc:.0%}, recommit {recommit_date})")

if __name__=="__main__": main()
