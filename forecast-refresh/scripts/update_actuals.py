#!/usr/bin/env python3
"""Overwrite the Actuals tab with real Logfood net rows and advance the as_of parameter.
Usage: update_actuals.py --xlsx WB --actuals CSV(date,dbu,net) --asof YYYY-MM-DD
The Math/Readout formulas + charts recompute automatically on open (fullCalcOnLoad)."""
import argparse, csv
from datetime import date, datetime
import openpyxl

def parse_d(s):
    s=s.strip()
    for f in ("%Y-%m-%d","%m/%d/%Y","%m/%d/%y"):
        try: return datetime.strptime(s,f).date()
        except ValueError: pass
    raise ValueError(f"bad date {s!r}")

def find_header(ws, needles):
    for r in range(1, 40):
        vals=[str(ws.cell(r,c).value or "").strip().lower() for c in range(1,12)]
        if all(any(n in v for v in vals) for n in needles): return r
    raise RuntimeError(f"header row not found for {needles}")

def col_of(ws, hr, name):
    for c in range(1,12):
        if name in str(ws.cell(hr,c).value or "").strip().lower(): return c
    raise RuntimeError(f"column {name!r} not found")

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--xlsx",required=True); ap.add_argument("--actuals",required=True)
    ap.add_argument("--asof",required=True)
    a=ap.parse_args()
    asof=parse_d(a.asof)
    rows=[]
    with open(a.actuals) as f:
        rd=csv.DictReader(f)
        for row in rd:
            k={kk.strip().lower():vv for kk,vv in row.items()}
            rows.append((parse_d(k["date"]), float(k["dbu"]), float(k["net"])))
    rows.sort()

    wb=openpyxl.load_workbook(a.xlsx)
    A=wb["Actuals"]
    hr=find_header(A,["date","dbu","logfood"])   # money col header is "$ (Logfood)"
    cD=col_of(A,hr,"date"); cB=col_of(A,hr,"dbu"); cN=col_of(A,hr,"logfood")
    cNote=cN+1
    # clear existing data rows
    r=hr+1
    while A.cell(r,cD).value not in (None,""):
        for c in range(cD,cNote+1): A.cell(r,c).value=None
        r+=1
    # write new rows
    from openpyxl.styles import Font
    f9=Font(name="DM Sans",size=9)
    for i,(dt,dbu,net) in enumerate(rows):
        rr=hr+1+i
        A.cell(rr,cD,dt).number_format="yyyy-mm-dd"; A.cell(rr,cD).font=f9
        A.cell(rr,cB,round(dbu,1)).number_format="#,##0"; A.cell(rr,cB).font=f9
        A.cell(rr,cN,round(net,0)).number_format="$#,##0"; A.cell(rr,cN).font=f9
        A.cell(rr,cNote,"logfood").font=Font(name="DM Sans",size=9,italic=True,color="595959")

    # advance as_of: find the labelled param cell on Event Ledger
    L=wb["Event Ledger"]
    set_it=False
    for r in range(1,20):
        for c in range(1,8):
            if "as_of" in str(L.cell(r,c).value or "").lower():
                L.cell(r,c+2, asof).number_format="yyyy-mm-dd"; set_it=True; break
        if set_it: break
    try: wb.calculation.fullCalcOnLoad=True
    except Exception: pass
    wb.save(a.xlsx)
    print(f"OK  wrote {len(rows)} actual rows ({rows[0][0]}..{rows[-1][0]})  as_of={asof}  set_asof={set_it}")

if __name__=="__main__": main()
