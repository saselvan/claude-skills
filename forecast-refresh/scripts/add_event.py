#!/usr/bin/env python3
"""Append a row to the Event Ledger. Usage:
add_event.py --xlsx WB --type contract|workload --date YYYY-MM-DD --status occurred|projected
  --label "..." --param <float> --confidence <0-1> --in-line Y|N --basis "..." [--id C9]"""
import argparse
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

def main():
    ap=argparse.ArgumentParser()
    for x in ["type","date","status","label","param","confidence","in-line","basis"]:
        ap.add_argument("--"+x,required=True)
    ap.add_argument("--xlsx",required=True); ap.add_argument("--id",default=None)
    a=ap.parse_args()
    wb=openpyxl.load_workbook(a.xlsx); L=wb["Event Ledger"]
    hr=None
    for r in range(1,30):
        if str(L.cell(r,2).value or "").strip().lower()=="event_id": hr=r; break
    r=hr+1
    while L.cell(r,2).value: r+=1
    eid=a.id or (("C" if a.type=="contract" else "W")+str(r-hr))
    dt=datetime.strptime(a.date,"%Y-%m-%d").date()
    vals=[eid,dt,a.type,a.status,a.label,float(a.param),float(a.confidence),getattr(a,"in_line"),a.basis]
    thin=Side(style="thin",color="D9D3C6"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
    for j,v in enumerate(vals):
        c=L.cell(r,2+j,v); c.border=box; c.font=Font(name="DM Sans",size=10)
        c.alignment=Alignment(horizontal="left",vertical="center",wrap_text=(j==8))
    L.cell(r,3).number_format="yyyy-mm-dd"; L.cell(r,7).number_format="0.0000"; L.cell(r,8).number_format="0%"
    L.cell(r,4).font=Font(name="DM Sans",size=10,bold=True,color=("C42B1A" if a.type=="contract" else "2272B4"))
    try: wb.calculation.fullCalcOnLoad=True
    except Exception: pass
    wb.save(a.xlsx)
    print(f"OK  added {eid}: {a.type} {a.status} {a.date} '{a.label}' param={a.param} in_line={a.in_line}")

if __name__=="__main__": main()
